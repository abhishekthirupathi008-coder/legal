import io
import os
import re
import sys
from datetime import date
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import html
import requests
import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF

BACKEND_URL = os.getenv("BACKEND_URL", "http://127.0.0.1:8000")
LOGO_PATH = PROJECT_ROOT / "Image" / "Logo.png"

st.set_page_config(page_title="LegalEase", page_icon="⚖️", layout="wide")

st.markdown("""
<style>
.main-title{font-size:42px;font-weight:800;margin-bottom:0}
.subtitle{font-size:18px;color:#666;margin-bottom:18px}
.preview{background:#fff;padding:34px;border:1px solid #ddd;border-radius:12px;line-height:1.7}
</style>
""", unsafe_allow_html=True)

if "document" not in st.session_state:
    st.session_state.document = ""
if "document_type" not in st.session_state:
    st.session_state.document_type = "NDA"

if LOGO_PATH.exists():
    a, b = st.columns([1, 6])
    with a:
        st.image(str(LOGO_PATH), width=90)
    with b:
        st.markdown('<div class="main-title">LegalEase</div>', unsafe_allow_html=True)
        st.markdown('<div class="subtitle">AI-Powered Legal Document Generator</div>', unsafe_allow_html=True)
else:
    st.markdown('<div class="main-title">⚖️ LegalEase</div>', unsafe_allow_html=True)
    st.markdown('<div class="subtitle">AI-Powered Legal Document Generator</div>', unsafe_allow_html=True)

st.warning("AI-generated legal drafts should be reviewed by a qualified legal professional before signing or relying on them.")

with st.sidebar:
    st.header("⚙️ Document Settings")
    document_type = st.selectbox(
        "Document Type",
        [
            "NDA", "Employment Contract", "Lease Agreement",
            "Employment Offer Letter", "Freelance Work Contract",
            "Service Agreement", "Partnership Agreement",
            "General Agreement", "Custom",
        ],
    )
    effective_date = st.date_input("Effective Date", value=date.today())
    st.divider()
    st.subheader("🏢 Branding")
    uploaded_logo = st.file_uploader("Upload Logo", type=["png", "jpg", "jpeg"])

st.subheader("📝 Document Information")
left, right = st.columns(2)
with left:
    parties = st.text_area(
        "Parties",
        height=180,
        placeholder="Example:\nSudhan - Disclosing Party\nABC Technologies - Receiving Party",
    )
with right:
    terms = st.text_area(
        "Terms and Requirements",
        height=180,
        placeholder="Example:\nConfidential information must remain private.\nAgreement duration is 2 years.\nTermination requires 15 days written notice.",
    )

if st.button("✨ Generate Legal Document", type="primary", use_container_width=True):
    if not parties.strip():
        st.error("Please enter the parties.")
    elif not terms.strip():
        st.error("Please enter the terms and requirements.")
    else:
        payload = {
            "document_type": document_type,
            "parties": parties,
            "terms": terms,
            "dates": str(effective_date),
        }
        with st.spinner("Generating document with Gemini AI..."):
            try:
                response = requests.post(f"{BACKEND_URL}/generate", json=payload, timeout=120)
                if response.ok:
                    data = response.json()
                    st.session_state.document = data["content"]
                    st.session_state.document_type = document_type
                    st.success("Document generated successfully! 🎉")
                else:
                    try:
                        detail = response.json().get("detail", response.text)
                    except ValueError:
                        detail = response.text
                    st.error(f"Generation failed: {detail}")
            except requests.exceptions.ConnectionError:
                st.error("Cannot connect to the FastAPI backend. Start it with: uvicorn legalEaseAPI.main:app --reload")
            except requests.exceptions.Timeout:
                st.error("The request timed out. Please try again.")
            except requests.RequestException as exc:
                st.error(f"Request error: {exc}")

if st.session_state.document:
    st.divider()
    st.header("📄 Generated Document")
    edit_tab, preview_tab = st.tabs(["✏️ Edit Document", "👁️ Preview"])

    with edit_tab:
        st.session_state.document = st.text_area(
            "Edit your document",
            value=st.session_state.document,
            height=650,
            key="document_editor",
        )

    with preview_tab:
        rendered = []
        for block in html.escape(st.session_state.document).split("\n\n"):
            for line in block.split("\n"):
                line = line.strip()
                if not line:
                    continue
                if line.isupper() or re.match(r"^\d+[.)]", line):
                    rendered.append(f"<h4>{line}</h4>")
                else:
                    rendered.append(f"<p>{line}</p>")
        st.markdown(f'<div class="preview">{"".join(rendered)}</div>', unsafe_allow_html=True)

    st.divider()
    st.subheader("⬇️ Export Document")
    document_text = st.session_state.document
    safe_name = re.sub(r"[^a-z0-9]+", "_", st.session_state.document_type.lower()).strip("_") or "legal_document"
    c1, c2, c3 = st.columns(3)

    with c1:
        st.download_button("📄 Download TXT", document_text, f"{safe_name}.txt", "text/plain", use_container_width=True)

    with c2:
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        if uploaded_logo:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(io.BytesIO(uploaded_logo.getvalue()), width=Inches(1.2))
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = title.add_run(st.session_state.document_type.upper())
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(16)
        for line in document_text.splitlines():
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(line)
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("LegalEase - AI-Assisted Draft")
        buf = io.BytesIO()
        doc.save(buf)
        st.download_button(
            "📝 Download DOCX", buf.getvalue(), f"{safe_name}.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )

    with c3:
        class LegalEasePDF(FPDF):
            def __init__(self, title):
                super().__init__()
                self.title_text = title
                self.set_auto_page_break(auto=True, margin=18)
            def header(self):
                self.set_font("Times", "B", 14)
                self.cell(0, 10, self.title_text.upper(), align="C")
                self.ln(10)
            def footer(self):
                self.set_y(-15)
                self.set_font("Times", "", 8)
                self.cell(0, 10, f"LegalEase - AI-Assisted Draft | Page {self.page_no()}", align="C")

        pdf = LegalEasePDF(st.session_state.document_type)
        pdf.set_margins(20, 20, 20)
        pdf.add_page()
        pdf.set_font("Times", "", 11)
        clean_text = (document_text.replace("—", "-").replace("–", "-")
                      .replace("“", '"').replace("”", '"')
                      .replace("‘", "'").replace("’", "'").replace("•", "-"))
        for line in clean_text.splitlines():
            line = line.strip()
            if not line:
                pdf.ln(4)
                continue
            try:
                line.encode("latin-1")
            except UnicodeEncodeError:
                line = line.encode("latin-1", "replace").decode("latin-1")
            pdf.set_font("Times", "B" if (line.isupper() or re.match(r"^\d+[.)]", line)) else "", 11)
            pdf.multi_cell(0, 6, line)
            pdf.ln(2)
        pdf_bytes = bytes(pdf.output())
        st.download_button("📕 Download PDF", pdf_bytes, f"{safe_name}.pdf", "application/pdf", use_container_width=True)

st.divider()
st.caption("LegalEase • AI-Powered Legal Document Generator")
